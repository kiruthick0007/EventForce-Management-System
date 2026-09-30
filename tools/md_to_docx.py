"""Convert PROJECT_REPORT.md into a submission-ready Word document.

The college specification requires documentation in Word or PDF, formatted in
Times New Roman 12 pt with clear headings. This script produces that .docx
directly, so the report stays maintainable as Markdown and is regenerated with
one command instead of being re-typed by hand.

Usage:  python tools/md_to_docx.py
"""

import io
import os
import re
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SOURCE = os.path.join(ROOT, "PROJECT_REPORT.md")
OUTPUT = os.path.join(ROOT, "EventForce_Project_Report.docx")

BODY_FONT = "Times New Roman"
CODE_FONT = "Consolas"

HEADING_SIZES = {"Heading 1": 20, "Heading 2": 16, "Heading 3": 14, "Heading 4": 12}

INLINE = re.compile(r"(\*\*\*.+?\*\*\*|\*\*.+?\*\*|\*[^*\n]+\*|`[^`\n]+`)")

# Filled in on the cover page; edit before submitting.
COVER_FIELDS = [
    ("Project Title", "EventForce \u2014 Event Management System"),
    ("Platform", "Salesforce (Lightning Experience)"),
    ("Submitted by", "[Your Name]"),
    ("Roll Number", "[Roll Number]"),
    ("Department", "[Department]"),
    ("Guided by", "[Guide Name]"),
    ("Institution", "[College Name]"),
    ("Academic Year", "[2025 \u2013 2026]"),
]


def set_style_font(style, name=BODY_FONT, size=None, bold=None, color=RGBColor(0, 0, 0)):
    style.font.name = name
    if size is not None:
        style.font.size = Pt(size)
    if bold is not None:
        style.font.bold = bold
    style.font.color.rgb = color
    rpr = style.element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.append(fonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(attr), name)


def shade(element, fill):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    element.append(shd)


def add_field(paragraph, instruction):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    text = OxmlElement("w:instrText")
    text.set(qn("xml:space"), "preserve")
    text.text = instruction
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for node in (begin, text, end):
        run._r.append(node)
    run.font.name = BODY_FONT
    run.font.size = Pt(9)


def add_runs(paragraph, text, bold=False, italic=False):
    """Write markdown inline formatting into a paragraph as real Word runs.

    Recursive, because the report nests code inside bold -- **`EventTrigger`**
    -- and a single pass would leave the backticks in the output as literal text.
    """
    for token in INLINE.split(text):
        if not token:
            continue
        for pattern, new_bold, new_italic in (
            (r"\*\*\*(.+?)\*\*\*", True, True),
            (r"\*\*(.+?)\*\*", True, None),
            (r"\*([^*\n]+)\*", None, True),
        ):
            match = re.fullmatch(pattern, token)
            if match:
                add_runs(paragraph, match.group(1),
                         bold if new_bold is None else new_bold,
                         italic if new_italic is None else new_italic)
                break
        else:
            code = re.fullmatch(r"`([^`\n]+)`", token)
            run = paragraph.add_run(code.group(1) if code else token)
            run.bold = bold
            run.italic = italic
            run.font.name = CODE_FONT if code else BODY_FONT
            run.font.size = Pt(10) if code else Pt(12)


def add_code_block(doc, lines):
    for line in lines:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.right_indent = Inches(0.25)
        shade(p._p.get_or_add_pPr(), "F2F2F2")
        run = p.add_run(line if line else " ")
        run.font.name = CODE_FONT
        run.font.size = Pt(9)
    doc.add_paragraph().paragraph_format.space_after = Pt(6)


def add_table(doc, rows):
    # rows[0] is the header row, rows[1] is the --- separator, rows[2:] the body.
    header = rows[0]
    body = rows[2:]
    if all(not c.strip() for c in header):
        # The metadata table at the top of the report has an empty header row.
        header = None

    columns = max(len(r) for r in rows)
    table = doc.add_table(rows=0, cols=columns)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    def emit(values, is_header):
        row = table.add_row()
        for index in range(columns):
            text = values[index] if index < len(values) else ""
            cell = row.cells[index]
            cell.text = ""
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_before = Pt(2)
            paragraph.paragraph_format.space_after = Pt(2)
            add_runs(paragraph, text, bold=is_header)
            if is_header:
                shade(cell._tc.get_or_add_tcPr(), "DCE6F1")

    if header is not None:
        emit(header, True)
    for values in body:
        emit(values, False)
    doc.add_paragraph().paragraph_format.space_after = Pt(6)


def split_row(line):
    return [c.strip() for c in line.strip().strip("|").split("|")]


def is_separator(line):
    return bool(re.match(r"^\s*\|?[\s:|-]+\|[\s:|-]*$", line)) and "-" in line


def add_rule(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)
    border = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "AAAAAA")
    border.append(bottom)
    p._p.get_or_add_pPr().append(border)


def add_cover(doc):
    for _ in range(3):
        doc.add_paragraph()
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("EventForce")
    run.font.name = BODY_FONT
    run.font.size = Pt(32)
    run.bold = True

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("Event Management System")
    run.font.name = BODY_FONT
    run.font.size = Pt(20)

    tagline = doc.add_paragraph()
    tagline.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = tagline.add_run("A Customer Relationship Management application on the Salesforce platform")
    run.font.name = BODY_FONT
    run.font.size = Pt(13)
    run.italic = True

    for _ in range(3):
        doc.add_paragraph()

    table = doc.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for label, value in COVER_FIELDS:
        row = table.add_row()
        for index, text in enumerate((label, value)):
            cell = row.cells[index]
            cell.text = ""
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_before = Pt(3)
            paragraph.paragraph_format.space_after = Pt(3)
            run = paragraph.add_run(text)
            run.font.name = BODY_FONT
            run.font.size = Pt(12)
            run.bold = index == 0

    note = doc.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = note.add_run("Replace the bracketed placeholders on this page before submission.")
    run.font.name = BODY_FONT
    run.font.size = Pt(10)
    run.italic = True

    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def build():
    doc = Document()

    section = doc.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    for attr in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
        setattr(section, attr, Inches(1))

    set_style_font(doc.styles["Normal"], size=12)
    doc.styles["Normal"].paragraph_format.space_after = Pt(8)
    doc.styles["Normal"].paragraph_format.line_spacing = 1.15
    for name, size in HEADING_SIZES.items():
        style = doc.styles[name]
        set_style_font(style, size=size, bold=True)
        style.paragraph_format.space_before = Pt(16 if name != "Heading 1" else 20)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.keep_with_next = True
    set_style_font(doc.styles["List Bullet"], size=12)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = header.add_run("EventForce \u2014 Salesforce Implementation Project")
    run.font.name = BODY_FONT
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_field(footer, "PAGE")
    run = footer.add_run(" of ")
    run.font.name = BODY_FONT
    run.font.size = Pt(9)
    add_field(footer, "NUMPAGES")

    add_cover(doc)

    with io.open(SOURCE, encoding="utf-8") as handle:
        lines = handle.read().split("\n")

    index = 0
    title_skipped = False
    while index < len(lines):
        line = lines[index]
        stripped = line.strip()

        if stripped.startswith("```"):
            block = []
            index += 1
            while index < len(lines) and not lines[index].strip().startswith("```"):
                block.append(lines[index])
                index += 1
            add_code_block(doc, block)
            index += 1
            continue

        if stripped.startswith("|") and index + 1 < len(lines) and is_separator(lines[index + 1]):
            rows = [split_row(stripped)]
            index += 1
            while index < len(lines) and lines[index].strip().startswith("|"):
                rows.append(split_row(lines[index].strip()))
                index += 1
            add_table(doc, rows)
            continue

        if not stripped:
            index += 1
            continue

        if stripped in ("---", "***", "___"):
            add_rule(doc)
            index += 1
            continue

        heading = re.match(r"^(#{1,6})\s+(.*)$", stripped)
        if heading:
            level = len(heading.group(1))
            text = heading.group(2)
            if level == 1 and not title_skipped:
                # The cover page already carries the title; drop it here and
                # promote every remaining heading one level so that "## Abstract"
                # becomes Heading 1 in the Word outline.
                title_skipped = True
                index += 1
                continue
            promoted = level - 1 if title_skipped else level
            paragraph = doc.add_heading("", level=max(1, min(promoted, 4)))
            paragraph.paragraph_format.keep_with_next = True
            add_runs(paragraph, text, bold=True)
            index += 1
            continue

        if stripped.startswith("> "):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.3)
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(4)
            shade(p._p.get_or_add_pPr(), "FFF7E6")
            add_runs(p, stripped[2:], italic=True)
            index += 1
            continue

        ordered = re.match(r"^(\d+)\.\s+(.*)$", stripped)
        if ordered:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.3)
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run(ordered.group(1) + ".  ")
            run.font.name = BODY_FONT
            run.font.size = Pt(12)
            add_runs(p, ordered.group(2))
            index += 1
            continue

        bullet = re.match(r"^[-*+]\s+(.*)$", stripped)
        if bullet:
            indent = len(line) - len(line.lstrip())
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.left_indent = Inches(0.3 + 0.25 * (indent // 2))
            p.paragraph_format.space_after = Pt(4)
            add_runs(p, bullet.group(1))
            index += 1
            continue

        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        add_runs(p, stripped)
        index += 1

    doc.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    path = build()
    sys.stdout.write("Wrote " + path + "\n")
